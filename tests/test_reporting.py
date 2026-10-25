from io import BytesIO
from openpyxl import load_workbook
from revenue_audit.analytics import snapshot
from revenue_audit.reporting import excel_bytes, build_reports


def test_workbooks_preserve_integer_cents(con, tmp_path, database):
    tables = snapshot(con)
    workbook = load_workbook(BytesIO(excel_bytes(tables)))
    assert "exceptions" in workbook.sheetnames
    sheet = workbook["reconciliation"]
    assert isinstance(sheet["A2"].value, int)
    assert sheet.freeze_panes == "A2"
    result = build_reports(database[0], tmp_path)
    assert result["invoice_count"] == 12024
    assert (tmp_path / "manifest.json").exists()
    assert (tmp_path / "risk_control_matrix.xlsx").exists()
